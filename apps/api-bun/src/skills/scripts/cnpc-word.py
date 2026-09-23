#!/usr/bin/env python3
"""生成或在既有 DOCX 文档末尾追加内容。

新建文档：
  {"title": "标题", "sections": [{"heading": "章节", "level": 1, "paragraphs": ["正文"]}]}

在既有文档末尾追加（source 为既有文件路径；既有文档已有标题，因此忽略顶层 title）：
  {"source": "既有.docx", "sections": [{"heading": "补充章节", "paragraphs": ["补充正文"]}]}
"""
import json
import sys
from docx import Document


def main():
    if len(sys.argv) != 3:
        raise SystemExit("用法: cnpc-word.py input.json output.docx")
    with open(sys.argv[1], encoding="utf-8") as stream:
        data = json.load(stream)
    source = data.get("source")
    if source:
        document = Document(source)
    else:
        document = Document()
        if data.get("title"):
            document.add_heading(str(data["title"]), 0)
    for section in data.get("sections", []):
        if section.get("heading"):
            document.add_heading(str(section["heading"]), level=min(max(int(section.get("level", 1)), 1), 9))
        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(str(paragraph))
        rows = section.get("table")
        if rows:
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = str(value)
    document.save(sys.argv[2])


if __name__ == "__main__":
    main()
