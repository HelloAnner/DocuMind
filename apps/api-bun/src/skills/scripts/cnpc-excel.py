#!/usr/bin/env python3
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font


def main():
    if len(sys.argv) != 3:
        raise SystemExit("用法: cnpc-excel.py input.json output.xlsx")
    with open(sys.argv[1], encoding="utf-8") as stream:
        data = json.load(stream)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_data in data.get("sheets", []):
        sheet = workbook.create_sheet(str(sheet_data.get("name", "Sheet"))[:31])
        for row in sheet_data.get("rows", []):
            sheet.append(list(row))
        if sheet.max_row:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        for column in sheet.columns:
            width = min(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 60)
            sheet.column_dimensions[column[0].column_letter].width = width
    if not workbook.sheetnames:
        workbook.create_sheet("Sheet1")
    workbook.save(sys.argv[2])


if __name__ == "__main__":
    main()
