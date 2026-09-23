#!/usr/bin/env python3
"""生成或修改 PPTX 演示文稿。

新建演示文稿：
  {"slides": [{"title": "标题", "bullets": ["要点一", "要点二"]}]}

在既有演示文稿末尾追加幻灯片（source 为既有文件路径）：
  {"source": "既有.pptx", "slides": [{"title": "标题", "bullets": ["要点一"]}]}
"""
import json
import sys
from pptx import Presentation


def main():
    if len(sys.argv) != 3:
        raise SystemExit("用法: cnpc-ppt.py input.json output.pptx")
    with open(sys.argv[1], encoding="utf-8") as stream:
        data = json.load(stream)
    source = data.get("source")
    if source:
        presentation = Presentation(source)
    else:
        presentation = Presentation()
    for slide_data in data.get("slides", []):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(slide_data.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(slide_data.get("bullets", [])):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
    if not presentation.slides:
        presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(sys.argv[2])


if __name__ == "__main__":
    main()
